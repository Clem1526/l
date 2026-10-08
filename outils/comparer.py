"""Compare les résultats de la macro (classeur exporté) au calcul de référence Python."""
import sys

import numpy as np
import openpyxl

sys.path.insert(0, __file__.rsplit("/", 2)[0] + "/reference")
import reference as ref  # noqa: E402

classeur, dossier = sys.argv[1], sys.argv[2]
ref.DOSSIER = dossier
wb = openpyxl.load_workbook(classeur, data_only=True)
T = ref.TITRES
data = {t: ref.lire(t) for t in T}
R = np.column_stack([ref.rendements(data[t][1], data[t][2]) for t in T])
mu, cov = R.mean(0), np.cov(R, rowvar=False, bias=True)
n = len(T)
ecarts = {}

st = wb["Statistiques"]
ecarts["E(R)"] = max(abs(st.cell(4, i + 2).value - mu[i]) for i in range(n))
ecarts["Var"] = max(abs(st.cell(5, i + 2).value - cov[i, i]) for i in range(n))
ecarts["Cov"] = max(abs(st.cell(13 + i, 2 + j).value - cov[i, j]) for i in range(n) for j in range(n))
ligne_verif = 13 + n + 2 + n + 4
ecarts["Controle Excel (formules)"] = max(st.cell(ligne_verif + 5, i + 2).value for i in range(n))
ecarts["COVARIANCE.P"] = max(abs(st.cell(ligne_verif + 8 + 1 + i, 2 + j).value - cov[i, j])
                             for i in range(n) for j in range(n))

rd = wb["Rendements"]
ecarts["Rendements"] = max(abs(rd.cell(t + 2, i + 2).value - R[t, i]) for t in range(len(R)) for i in range(n))

dt = wb["Deux titres"]
a = T.index(dt["B3"].value.split("(")[1].rstrip(")"))
b = T.index(dt["C3"].value.split("(")[1].rstrip(")"))
for k in range(11):
    xa = dt.cell(11 + k, 1).value
    w = np.zeros(n); w[a], w[b] = xa, 1 - xa
    ecarts.setdefault("Deux titres", 0)
    ecarts["Deux titres"] = max(ecarts["Deux titres"], abs(dt.cell(11 + k, 3).value - mu @ w),
                                abs(dt.cell(11 + k, 5).value - np.sqrt(w @ cov @ w)))

fr = wb["Frontiere"]
ligne, pire_sigma, pire_scipy = 9, 0, 0
while fr.cell(ligne, 1).value is not None:
    w = np.array([fr.cell(ligne, 8 + i).value for i in range(n)])
    r_pf, s_pf = fr.cell(ligne, 3).value, fr.cell(ligne, 5).value
    pire_sigma = max(pire_sigma, abs(s_pf - np.sqrt(w @ cov @ w)), abs(r_pf - mu @ w), abs(w.sum() - 1))
    ws_ = ref.qp_scipy(cov, mu, r_pf)
    pire_scipy = max(pire_scipy, np.sqrt(w @ cov @ w) - np.sqrt(ws_ @ cov @ ws_))
    ligne += 1
ecarts["Frontiere (coherence poids/R/sigma)"] = pire_sigma
ecarts["Frontiere : sigma macro - sigma optimal scipy"] = pire_scipy
print(f"{ligne - 9} points sur la frontière")

ca = wb["Calculateur"]
w = np.ones(n) / n
ecarts["Calculateur R"] = abs(ca.cell(6 + n + 2, 2).value - mu @ w)
ecarts["Calculateur sigma"] = abs(ca.cell(6 + n + 4, 2).value - np.sqrt(w @ cov @ w))
print("Contrôle calculateur :", ca.cell(6 + n + 7, 2).value)

for k, v in ecarts.items():
    print(f"{k:50s} {v:.2e}  {'OK' if v < 1e-9 else 'A VERIFIER'}")
