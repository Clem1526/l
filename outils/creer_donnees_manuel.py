"""Crée Donnees_2_titres.xlsx : cours et dividendes de deux actions, prêts pour la version manuelle."""
import csv
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

dossier, sortie, a, b = sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4]
lire = lambda t: list(csv.DictReader(open(f"{dossier}/{t}.csv")))
ra, rb = lire(a), lire(b)
assert [r["Date"] for r in ra] == [r["Date"] for r in rb]
wb = Workbook()
ws = wb.active
ws.title = "Donnees"
entetes = ["Date", f"Cours {a}", f"Dividende {a}", f"Cours {b}", f"Dividende {b}"]
for j, e in enumerate(entetes, 1):
    c = ws.cell(1, j, e)
    c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="002060")
for i, (x, y) in enumerate(zip(ra, rb), 2):
    ws.cell(i, 1, date.fromisoformat(x["Date"])).number_format = "dd/mm/yyyy"
    for j, v in enumerate([x["Close"], x["Dividend"], y["Close"], y["Dividend"]], 2):
        ws.cell(i, j, float(v)).number_format = "0.00"
for col in "ABCDE":
    ws.column_dimensions[col].width = 14
wb.save(sortie)
print("Créé :", sortie, len(ra), "lignes")
