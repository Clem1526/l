"""Génère le classeur modèle de la version simplifiée (Mode d'emploi, Parametres, Dividendes)."""
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SORTIE = sys.argv[1]
BLEU, JAUNE, VERT, OR = "1F3864", "FFF2CC", "00823C", "FFD966"
fond = lambda c: PatternFill("solid", fgColor=c)
blanc_gras = Font(bold=True, color="FFFFFF")

wb = Workbook()
# ------------------------------------------------------------------ Mode d'emploi
me = wb.active
me.title = "Mode d'emploi"
me.sheet_properties.tabColor = BLEU
me.column_dimensions["A"].width = 3
me.column_dimensions["B"].width = 115
blocs = [
    ("t", "Frontière efficiente — mode d'emploi"),
    ("", ""),
    ("h", "Ce que fait la macro"),
    ("", "1. Télécharge sur Euronext les cours journaliers des actions choisies (et ajoute les dividendes de la feuille Dividendes)."),
    ("", "2. Calcule pour chaque action le rendement moyen, la variance et l'écart-type."),
    ("", "3. Calcule la covariance (et la corrélation) entre les actions."),
    ("", "4. Calcule le rendement et le risque de TOUS les portefeuilles possibles (poids par pas de 5 %)."),
    ("", "5. Garde, pour chaque niveau de rendement, le portefeuille le moins risqué : c'est la frontière efficiente (tableau + graphique)."),
    ("", ""),
    ("h", "Utilisation"),
    ("", "• Feuille Parametres : choisissez les actions (Oui / Non), leur code ISIN (sur live.euronext.com), la période et les réglages."),
    ("", "• Cliquez sur « Lancer l'analyse ». Un message résume les résultats ; en cas de problème, il indique quoi corriger."),
    ("", "• Sans connexion : choisissez la source « Fichiers CSV » et placez à côté du classeur un export Euronext par action (MC.csv…)."),
    ("", ""),
    ("h", "Les feuilles de résultats (couleur de l'onglet)"),
    ("", "• Cours (bleu) : cours de clôture et rendements R(t) = (D(t) + P(t) − P(t−1)) / P(t−1)."),
    ("", "• Statistiques (vert) : E(R), Var(R), σ(R), matrice de covariance, corrélations colorées, contrôle avec les fonctions Excel."),
    ("", "• Deux titres (orange) : le tableau xA / xB → R(pf), var(pf), σ(pf) et son graphique ; en doré, le portefeuille le moins risqué."),
    ("", "• Frontiere (doré) : un portefeuille par niveau de rendement ; vert = efficient, gris = dominé, doré = variance minimale."),
    ("", "• Portefeuilles (gris) : la liste de tous les portefeuilles testés (le nuage gris du graphique)."),
    ("", ""),
    ("h", "Hypothèses (cours « Utiliser la théorie du portefeuille »)"),
    ("", "• Données journalières ; variance, écart-type et covariance divisés par n."),
    ("", "• Pas de vente à découvert : chaque poids est entre 0 et 100 % et la somme fait 100 %."),
    ("", "• Euronext ne fournit que les deux dernières années de cours et pas les dividendes (d'où la feuille Dividendes)."),
    ("", ""),
    ("h", "Installation (une seule fois)"),
    ("", "1. Enregistrez ce classeur au format « Classeur Excel prenant en charge les macros » (.xlsm)."),
    ("", "2. Outils > Macro > Éditeur Visual Basic, Insertion > Module, collez le contenu de CODE_A_COPIER.txt."),
    ("", "3. Outils > Macro > Macros… > CreerBoutons > Exécuter : les boutons apparaissent sur la feuille Parametres."),
]
for i, (style, texte) in enumerate(blocs, start=1):
    c = me.cell(i, 2, texte)
    if style == "t":
        c.font = Font(bold=True, size=16, color=BLEU)
    elif style == "h":
        c.font, c.fill = blanc_gras, fond(BLEU)

# ------------------------------------------------------------------ Parametres
pa = wb.create_sheet("Parametres")
pa.sheet_properties.tabColor = VERT
for col, w in zip("ABCDEF", (3, 46, 18, 20, 18, 4)):
    pa.column_dimensions[col].width = w
pa["B1"] = "Paramètres de l'analyse"
pa["B1"].font = Font(bold=True, size=16, color=BLEU)
for c, t in (("B3", "Réglage"), ("C3", "Valeur")):
    pa[c] = t
    pa[c].font, pa[c].fill = blanc_gras, fond(BLEU)
reglages = [
    ("Source des cours", "Euronext", None),
    ("Date de début", date(2025, 1, 1), "dd/mm/yyyy"),
    ("Date de fin", date(2025, 12, 31), "dd/mm/yyyy"),
    ("Pas des poids testés (5 % conseillé)", 0.05, "0%"),
    ("Nombre de points de la frontière", 20, "0"),
    ("Pas du tableau à deux titres", 0.10, "0%"),
    ("Titre A du tableau à deux titres", "BNP", None),
    ("Titre B du tableau à deux titres", "AI", None),
]
for i, (lib, val, fmt) in enumerate(reglages, start=4):
    pa.cell(i, 2, lib)
    c = pa.cell(i, 3, val)
    c.fill, c.alignment = fond(JAUNE), Alignment(horizontal="center")
    if fmt:
        c.number_format = fmt
dv = DataValidation(type="list", formula1='"Euronext,Fichiers CSV"')
pa.add_data_validation(dv)
dv.add("C4")
pa["B13"] = "Actions à analyser"
pa["B13"].font = Font(bold=True, size=12, color=BLEU)
for col, t in zip("BCDE", ("Utiliser ? (Oui / Non)", "Code", "Nom de l'entreprise", "Code ISIN")):
    pa[f"{col}14"] = t
    pa[f"{col}14"].font, pa[f"{col}14"].fill = blanc_gras, fond(BLEU)
    pa[f"{col}14"].alignment = Alignment(horizontal="center")
actions = [("MC", "LVMH", "FR0000121014"), ("TTE", "TotalEnergies", "FR0000120271"),
           ("SAN", "Sanofi", "FR0000120578"), ("BNP", "BNP Paribas", "FR0000131104"),
           ("AI", "Air Liquide", "FR0000120073")]
for r in range(15, 45):
    for col in "BCDE":
        pa[f"{col}{r}"].fill = fond(JAUNE)
    pa[f"B{r}"].alignment = Alignment(horizontal="center")
for r, (code, nom, isin) in enumerate(actions, start=15):
    pa[f"B{r}"], pa[f"C{r}"], pa[f"D{r}"], pa[f"E{r}"] = "Oui", code, nom, isin
dv2 = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True)
pa.add_data_validation(dv2)
dv2.add("B15:B44")
pa["G3"] = "Boutons : lancer une fois la macro CreerBoutons"
pa["G3"].font = Font(italic=True, color="808080")

# ------------------------------------------------------------------ Dividendes
di = wb.create_sheet("Dividendes")
di.sheet_properties.tabColor = OR
for col, w in zip("ABCD", (10, 20, 20, 50)):
    di.column_dimensions[col].width = w
di["A1"] = "Dividendes détachés (Euronext ne les fournit pas)"
di["A1"].font = Font(bold=True, size=14, color=BLEU)
di["A2"] = "Ajoutés au rendement du jour de détachement : R(t) = (D(t) + P(t) − P(t−1)) / P(t−1)."
di["A2"].font = Font(italic=True, color="808080")
for col, t in zip("ABCD", ("Code", "Date de détachement", "Montant par action (€)", "Commentaire")):
    di[f"{col}3"] = t
    di[f"{col}3"].font, di[f"{col}3"].fill = blanc_gras, fond(BLEU)
divs = [("TTE", date(2025, 1, 2), 0.79, "Acompte (1er jour : sans effet)"), ("TTE", date(2025, 3, 26), 0.79, "Acompte"),
        ("MC", date(2025, 4, 24), 7.50, "Solde 2024"), ("SAN", date(2025, 5, 12), 3.92, "Exercice 2024"),
        ("BNP", date(2025, 5, 19), 4.79, "Solde 2024"), ("AI", date(2025, 5, 19), 3.30, "Exercice 2024"),
        ("TTE", date(2025, 6, 19), 0.85, "Acompte"), ("BNP", date(2025, 9, 26), 2.59, "Acompte 2025"),
        ("TTE", date(2025, 10, 1), 0.85, "Acompte"), ("MC", date(2025, 12, 2), 5.50, "Acompte 2025"),
        ("TTE", date(2025, 12, 31), 0.85, "Acompte")]
for r, (code, d, m, com) in enumerate(divs, start=4):
    di.cell(r, 1, code)
    di.cell(r, 2, d).number_format = "dd/mm/yyyy"
    di.cell(r, 3, m).number_format = "0.00"
    di.cell(r, 4, com)
    for c in (1, 2, 3):
        di.cell(r, c).fill = fond(JAUNE)
wb.active = 1
wb.save(SORTIE)
print("Créé :", SORTIE)
